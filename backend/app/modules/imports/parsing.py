"""Parseo de ficheros de importación: CSV (`;`/`,`), XLSX y XLS. Ref: RF-01, RF-04, RNF-09."""

from __future__ import annotations

import csv
import io
import zipfile

import xlrd
from openpyxl import load_workbook
from xlrd.compdoc import CompDocError

from app.core.errors import DomainError

XLSX_MAGIC = b"PK\x03\x04"
# Compound File Binary / OLE2 (Excel 97-2003). Los 4 primeros bytes bastan; el resto
# del identificador OLE es A1 B1 1A E1.
XLS_MAGIC = b"\xd0\xcf\x11\xe0"
XLS_VBA_PROJECT_UTF16 = "_VBA_PROJECT".encode("utf-16-le")
MAX_ROWS = 500

DEFAULT_COLUMN_MAPPING = {
    "id_paciente": "id_paciente",
    "nombre_referencia": "nombre_referencia",
    "direccion": "direccion",
    "codigo_postal": "codigo_postal",
    "municipio": "municipio",
    "provincia": "provincia",
    "tipo_zona": "tipo_zona",
}


def detect_format(*, filename: str, content: bytes) -> str:
    """Valida la firma real del fichero (no solo la extensión). 415 si no coincide."""
    lower_name = filename.lower()
    if lower_name.endswith(".xlsx"):
        if not content.startswith(XLSX_MAGIC):
            raise DomainError(
                415, "INVALID_FILE_SIGNATURE", "El fichero .xlsx no tiene una firma válida"
            )
        _reject_xlsx_macros(content)
        return "xlsx"
    if lower_name.endswith(".xls"):
        if not content.startswith(XLS_MAGIC):
            raise DomainError(
                415, "INVALID_FILE_SIGNATURE", "El fichero .xls no tiene una firma válida"
            )
        _reject_xls_macros(content)
        return "xls"
    if lower_name.endswith(".csv"):
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DomainError(
                415, "INVALID_FILE_SIGNATURE", "El fichero .csv no es texto UTF-8 válido"
            ) from exc
        return "csv"
    raise DomainError(
        415, "UNSUPPORTED_FORMAT", "Formato de fichero no soportado (usa .csv, .xlsx o .xls)"
    )


def _detect_csv_delimiter(sample: str) -> str:
    first_line = sample.splitlines()[0] if sample else ""
    return ";" if first_line.count(";") >= first_line.count(",") else ","


def _parse_csv(content: bytes) -> list[dict[str, str | None]]:
    text = content.decode("utf-8-sig")
    delimiter = _detect_csv_delimiter(text)
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return list(reader)


def _reject_xlsx_macros(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile:
        return
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise DomainError(
            422, "MACROS_NOT_ALLOWED", "El fichero contiene macros y no está permitido"
        )


def _reject_xls_macros(content: bytes) -> None:
    if b"_VBA_PROJECT" in content or XLS_VBA_PROJECT_UTF16 in content:
        raise DomainError(
            422, "MACROS_NOT_ALLOWED", "El fichero contiene macros y no está permitido"
        )


def _cell_text(value: object) -> str | None:
    if value is None or value == "":
        return None
    return str(value)


def _parse_xlsx(content: bytes) -> list[dict[str, str | None]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in (next(rows_iter, None) or [])]
        return [
            {header: _cell_text(value) for header, value in zip(headers, row, strict=False)}
            for row in rows_iter
            if any(cell is not None for cell in row)
        ]
    finally:
        workbook.close()


def _parse_xls(content: bytes) -> list[dict[str, str | None]]:
    try:
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
    except (xlrd.XLRDError, CompDocError, OSError) as exc:
        raise DomainError(
            415, "INVALID_FILE_SIGNATURE", "El fichero .xls no es un Excel binario válido"
        ) from exc
    try:
        if book.nsheets < 1:
            return []
        sheet = book.sheet_by_index(0)
        if sheet.nrows == 0:
            return []
        headers = [
            (_cell_text(sheet.cell_value(0, col)) or "").strip() for col in range(sheet.ncols)
        ]
        rows: list[dict[str, str | None]] = []
        for rowx in range(1, sheet.nrows):
            values = [sheet.cell_value(rowx, colx) for colx in range(sheet.ncols)]
            if not any(value not in ("", None) for value in values):
                continue
            rows.append(
                {header: _cell_text(value) for header, value in zip(headers, values, strict=False)}
            )
        return rows
    finally:
        book.release_resources()


def parse_rows(
    *, filename: str, content: bytes, mapping: dict[str, str] | None = None
) -> list[dict[str, str | None]]:
    """Devuelve filas con claves canónicas (`id_paciente`, `direccion`, ...)."""
    fmt = detect_format(filename=filename, content=content)
    column_mapping = mapping or DEFAULT_COLUMN_MAPPING

    if fmt == "csv":
        raw_rows = _parse_csv(content)
    elif fmt == "xls":
        raw_rows = _parse_xls(content)
    else:
        raw_rows = _parse_xlsx(content)

    if len(raw_rows) > MAX_ROWS:
        raise DomainError(
            422, "IMPORT_TOO_MANY_ROWS", f"El fichero supera el máximo de {MAX_ROWS} filas"
        )

    return [
        {
            canonical: raw_row.get(source)
            for canonical, source in column_mapping.items()
            if source in raw_row
        }
        for raw_row in raw_rows
    ]
